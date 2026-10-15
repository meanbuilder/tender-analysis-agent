from io import BytesIO

from openpyxl import load_workbook

from src.boq_extractor import BOQItem
from src.excel_exporter import (
    create_tender_workbook,
    export_tender_workbook,
    workbook_to_bytes,
)
from src.schemas import (
    EvidenceRecord,
    FindingStatus,
    ParameterFinding,
    RiskLevel,
    TenderAnalysisResult,
)


def sample_result() -> TenderAnalysisResult:
    payment = ParameterFinding(
        parameter_number=6,
        parameter_key="payment_terms",
        parameter_name="Payment Terms",
        status=FindingStatus.FOUND,
        requirement_summary="Payment within 30 days of receipt.",
        costing_impact="High",
        risk_level=RiskLevel.MEDIUM,
        evidence=[
            EvidenceRecord(
                filename="ATC.pdf",
                page_number=8,
                quoted_text="Payment within 30 days of receipt.",
                clause_reference="Clause 4.1",
            )
        ],
        review_required=True,
        review_reason="Verify the complete payment clause.",
    )

    return TenderAnalysisResult(
        tender_reference="BID-TEST-001",
        findings=[payment],
    )


def test_workbook_has_expected_five_sheets():
    workbook = create_tender_workbook(sample_result())

    assert workbook.sheetnames == [
        "Tender Summary",
        "BOQ and Technical Items",
        "Compliance and Risk Register",
        "Evidence and Exceptions",
        "Make in India - Annexure 8",
    ]


def test_summary_preserves_evidence_reference():
    workbook = create_tender_workbook(sample_result())
    ws = workbook["Tender Summary"]

    assert ws.cell(2, 2).value == "Payment Terms"
    assert ws.cell(2, 7).value == "ATC.pdf, page 8"
    assert ws.cell(2, 9).value == "Yes"


def test_evidence_sheet_contains_source_excerpt():
    workbook = create_tender_workbook(sample_result())
    ws = workbook["Evidence and Exceptions"]

    assert ws.cell(2, 2).value == "ATC.pdf"
    assert ws.cell(2, 3).value == 8
    assert ws.cell(2, 5).value == "Payment within 30 days of receipt."


def test_boq_sheet_exports_candidate_items():
    item = BOQItem(
        item_number="1",
        description="Industrial Valves",
        quantity=10,
        unit="Nos",
        source_filename="BOQ.pdf",
        source_page=3,
        review_required=True,
    )

    workbook = create_tender_workbook(sample_result(), [item])
    ws = workbook["BOQ and Technical Items"]

    assert ws.cell(2, 2).value == "Industrial Valves"
    assert ws.cell(2, 3).value == 10
    assert ws.cell(2, 7).value == "BOQ.pdf"
    assert ws.cell(2, 8).value == 3


def test_annexure_does_not_invent_local_content():
    workbook = create_tender_workbook(sample_result())
    ws = workbook["Make in India - Annexure 8"]

    values = [ws.cell(row, 2).value for row in range(2, ws.max_row + 1)]

    assert any("Not calculated" in value for value in values)
    assert any("authorized" in value.lower() for value in values)


def test_workbook_serializes_to_valid_xlsx_bytes():
    data = workbook_to_bytes(sample_result())

    assert data[:2] == b"PK"

    reopened = load_workbook(BytesIO(data), read_only=True)
    assert "Tender Summary" in reopened.sheetnames
    reopened.close()


def test_export_creates_file(tmp_path):
    output = tmp_path / "reports" / "tender_analysis.xlsx"

    returned_path = export_tender_workbook(sample_result(), output)

    assert returned_path.exists()
    assert returned_path.stat().st_size > 0

    workbook = load_workbook(returned_path, read_only=True)
    assert len(workbook.sheetnames) == 5
    workbook.close()


from openpyxl import Workbook

from src.excel_exporter import (
    _append_excel_row,
    _sanitize_excel_value,
)


def test_sanitize_excel_value_removes_illegal_control_characters():
    text = "Bid\x01 number\x0b: NTPC\x1f"
    assert _sanitize_excel_value(text) == "Bid number: NTPC"


def test_sanitize_excel_value_preserves_hindi_and_english():
    text = "Bid: NTPC | न्यूनतम बोलियों की संख्या"
    assert _sanitize_excel_value(text) == text


def test_sanitize_excel_value_preserves_valid_whitespace():
    text = "Line 1\nLine 2\tColumn"
    assert _sanitize_excel_value(text) == text


def test_append_excel_row_handles_illegal_characters_and_unicode():
    workbook = Workbook()
    worksheet = workbook.active

    _append_excel_row(
        worksheet,
        [
            "NTPC\x01 Tender",
            "बिड लगाने की समय सीमा",
            "Line 1\nLine 2",
        ],
    )

    assert worksheet["A1"].value == "NTPC Tender"
    assert worksheet["B1"].value == "बिड लगाने की समय सीमा"
    assert worksheet["C1"].value == "Line 1\nLine 2"
